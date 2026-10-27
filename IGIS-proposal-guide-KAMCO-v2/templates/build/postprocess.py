"""Finishing pass for the pptxgenjs output (build-template.js).

pptxgenjs cannot express a few things the IGIS spec needs, so they are
patched into the package here:

  1. Theme fonts — East Asian + Hangul script slots → Pretendard SemiBold
     (headings) / Pretendard (body). pptxgenjs only sets the Latin slot, so
     Korean would otherwise fall back to 맑은 고딕.
     Theme format scheme flattened: solid fills only, no gradients, no
     shadows, so PowerPoint's shape-style presets can't add effects that
     print differently from the screen.
  2. Two-panel bullets — ● bullets on the panel body placeholders,
     GRAY-400 on panel A and BLUE on panel B (layout + sample slide).
  3. Waterfall chart — hide the invisible base series and zero labels,
     signed delta labels (+38 / −22), white labels on BLUE / NAVY bars.
  4. --potx: re-type the package as a PowerPoint template.

Usage: python3 postprocess.py deck.pptx [--potx out.potx]
"""

import re
import shutil
import sys
import zipfile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_LABEL_POSITION
from pptx.util import Pt

HEAD_FONT, BODY_FONT = "Pretendard SemiBold", "Pretendard"
GRAY, G400, WHITE, NAVY = "5A5A5A", "A8A8A8", "FFFFFF", "2D375A"
BLUE = "0055A5"

BULLET_PPR = (
    '<a:lvl1pPr marL="228600" indent="-228600">'
    '<a:spcAft><a:spcPts val="1200"/></a:spcAft>'
    '<a:buClr><a:srgbClr val="{color}"/></a:buClr>'
    '<a:buSzPct val="70000"/><a:buFont typeface="Pretendard"/><a:buChar char="&#x25CF;"/>'
    "</a:lvl1pPr>"
)
PANEL_BULLETS = {"panel A body": G400, "panel B body": BLUE}

# Waterfall series → (number format, label color). Format sections hide zeros.
WATERFALL = {
    "기초": ('0;;;', GRAY),
    "증가": ('"+"0;;;', GRAY),
    "감소": ('"−"0;;;', GRAY),
    "강조": ('"+"0;;;', WHITE),
    "기말": ('0;;;', WHITE),
}


def patch_theme(xml):
    def fill(block, font):
        block = re.sub(r'<a:ea typeface="[^"]*"/>', f'<a:ea typeface="{font}"/>', block)
        return re.sub(r'<a:font script="Hang" typeface="[^"]*"/>', f'<a:font script="Hang" typeface="{font}"/>', block)

    xml = re.sub(r"<a:majorFont>.*?</a:majorFont>", lambda m: fill(m.group(0), HEAD_FONT), xml, flags=re.S)
    xml = re.sub(r"<a:minorFont>.*?</a:minorFont>", lambda m: fill(m.group(0), BODY_FONT), xml, flags=re.S)
    solid = '<a:solidFill><a:schemeClr val="phClr"/></a:solidFill>' * 3
    xml = re.sub(r"<a:fillStyleLst>.*?</a:fillStyleLst>", f"<a:fillStyleLst>{solid}</a:fillStyleLst>", xml, flags=re.S)
    xml = re.sub(r"<a:bgFillStyleLst>.*?</a:bgFillStyleLst>", f"<a:bgFillStyleLst>{solid}</a:bgFillStyleLst>", xml, flags=re.S)
    flat = "<a:effectStyle><a:effectLst/></a:effectStyle>" * 3
    return re.sub(r"<a:effectStyleLst>.*?</a:effectStyleLst>", f"<a:effectStyleLst>{flat}</a:effectStyleLst>", xml, flags=re.S)


def sp_blocks(xml):
    return re.finditer(r"<p:sp>.*?</p:sp>", xml, flags=re.S)


def patch_layout(xml, panel_idx):
    """Give the panel body placeholders a bulleted lvl1 style; record their idx."""
    for m in list(sp_blocks(xml)):
        block = m.group(0)
        name = re.search(r'<p:cNvPr id="\d+" name="([^"]*)"', block).group(1)
        if name not in PANEL_BULLETS:
            continue
        panel_idx[re.search(r'idx="(\d+)"', block).group(1)] = PANEL_BULLETS[name]
        style = "<a:lstStyle>" + BULLET_PPR.format(color=PANEL_BULLETS[name]) + "</a:lstStyle>"
        new = re.sub(r"<a:lstStyle\s*/>|<a:lstStyle>.*?</a:lstStyle>", style, block, count=1, flags=re.S)
        xml = xml.replace(block, new)
    return xml


def patch_slide(xml, panel_idx):
    """Slide-level paragraphs in panel bodies: drop pptxgenjs' buNone so the layout bullets show."""
    for m in list(sp_blocks(xml)):
        block = m.group(0)
        idx = re.search(r'<p:ph[^>]*idx="(\d+)"', block)
        if not idx or idx.group(1) not in panel_idx:
            continue
        new = block.replace("<a:buNone/>", "").replace(' indent="0" marL="0"', "")
        xml = xml.replace(block, new)
    return xml


def rewrite_package(path, potx_out=None):
    tmp = path + ".tmp"
    panel_idx = {}
    with zipfile.ZipFile(path) as zin:
        names = zin.namelist()
        # layouts first so panel idx values are known when slides are patched
        order = sorted(names, key=lambda n: 0 if "slideLayouts/slideLayout" in n else 1)
        out = {}
        for n in order:
            data = zin.read(n)
            if n.startswith("ppt/theme/") and n.endswith(".xml"):
                data = patch_theme(data.decode()).encode()
            elif re.match(r"ppt/slideLayouts/slideLayout\d+\.xml$", n):
                data = patch_layout(data.decode(), panel_idx).encode()
            elif re.match(r"ppt/slides/slide\d+\.xml$", n):
                data = patch_slide(data.decode(), panel_idx).encode()
            elif n == "[Content_Types].xml" and potx_out:
                # pptxgenjs only declares the master when a slide exists
                master = (b'<Override PartName="/ppt/slideMasters/slideMaster1.xml" ContentType='
                          b'"application/vnd.openxmlformats-officedocument.presentationml.slideMaster+xml"/>')
                if b"/ppt/slideMasters/slideMaster1.xml" not in data:
                    data = data.replace(b"</Types>", master + b"</Types>")
                data = data.replace(
                    b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
                    b"application/vnd.openxmlformats-officedocument.presentationml.template.main+xml",
                )
            out[n] = data
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:  # keep original order ([Content_Types].xml first)
            zout.writestr(n, out[n])
    shutil.move(tmp, potx_out or path)


def patch_charts(path):
    prs = Presentation(path)
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_chart:
                continue
            plot = shape.chart.plots[0]
            for series in plot.series:
                if series.name == "base":
                    series.format.fill.background()
                    series.format.line.fill.background()
                    dl = series.data_labels
                    dl.show_value = False
                    continue
                if series.name not in WATERFALL:
                    continue
                fmt, color = WATERFALL[series.name]
                dl = series.data_labels
                dl.number_format = fmt
                dl.number_format_is_linked = False
                dl.show_value = True
                dl.position = XL_LABEL_POSITION.INSIDE_END
                dl.font.size = Pt(12)
                dl.font.bold = False  # SemiBold is its own family; never fake-bold it
                dl.font.name = HEAD_FONT
                dl.font.color.rgb = RGBColor.from_string(color)
            ticks = shape.chart.category_axis.tick_labels
            ticks.font.name = BODY_FONT
    prs.save(path)


if __name__ == "__main__":
    src = sys.argv[1]
    potx = sys.argv[sys.argv.index("--potx") + 1] if "--potx" in sys.argv else None
    patch_charts(src)
    rewrite_package(src, potx)
    if potx:
        import os
        os.remove(src)
    print("post-processed", potx or src)
